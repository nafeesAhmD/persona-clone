import io
import uuid

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import Document
from app.services.document_extractor import normalize_text
from tests.conftest import USER

CLONE = {"name": "Ayaan AI", "bio": "Developer", "owner_consent": True}


# --- Test files ---------------------------------------------------------------


def make_pdf(text: str | None) -> bytes:
    """A minimal valid one-page PDF. With text=None the page has no text layer (like a scan)."""
    content = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return pdf


def make_docx(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "FastAPI"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- Fixtures -----------------------------------------------------------------


@pytest.fixture
def owner(client: TestClient, registered: dict) -> TestClient:
    return client


@pytest.fixture
def clone_id(owner: TestClient) -> str:
    response = owner.post("/api/clones", json=CLONE)
    assert response.status_code == 201
    return response.json()["id"]


@pytest.fixture
def other(make_client) -> TestClient:
    browser = make_client()
    assert browser.post(
        "/api/auth/register", json={**USER, "email": "sara@example.com", "username": "sara"}
    ).status_code == 201
    return browser


def upload(client: TestClient, clone_id: str, name: str, data: bytes, mime: str = "application/octet-stream", **form):
    return client.post(f"/api/clones/{clone_id}/documents", files={"file": (name, data, mime)}, data=form)


def stored_files(storage_dir) -> list:
    return [p for p in storage_dir.rglob("*") if p.is_file()]


# --- Uploads ------------------------------------------------------------------


def test_upload_txt(owner: TestClient, clone_id: str, storage_dir, qdrant):
    response = upload(owner, clone_id, "notes.txt", "Hello   world.\r\n\r\n\r\n\r\nSecond  paragraph.".encode(), "text/plain")
    assert response.status_code == 201, response.text
    doc = response.json()
    assert doc["status"] == "ready"
    assert doc["source_type"] == "file"
    assert doc["file_type"] == "txt"
    assert doc["title"] == "notes"  # defaults to the filename without extension
    assert doc["character_count"] == len("Hello world.\n\nSecond paragraph.")
    assert "stored_filename" not in doc and "extracted_text" not in doc  # no internals, no big text
    assert len(stored_files(storage_dir)) == 1
    # Ready text is indexed into AI memory automatically.
    assert doc["indexing_status"] == "indexed" and doc["chunk_count"] == 1
    assert qdrant.count_chunks(clone_id=uuid.UUID(clone_id), document_id=uuid.UUID(doc["id"])) == 1


def test_upload_pdf(owner: TestClient, clone_id: str):
    response = upload(owner, clone_id, "cv.pdf", make_pdf("AWS Certified Developer 2024"), "application/pdf", title="My CV")
    assert response.status_code == 201
    doc = response.json()
    assert (doc["status"], doc["file_type"], doc["title"], doc["mime_type"]) == ("ready", "pdf", "My CV", "application/pdf")
    detail = owner.get(f"/api/clones/{clone_id}/documents/{doc['id']}").json()
    assert "AWS Certified Developer 2024" in detail["extracted_text"]


def test_upload_docx(owner: TestClient, clone_id: str):
    data = make_docx(["I started with Python.", "", "Then Django and FastAPI."])
    response = upload(owner, clone_id, "journey.docx", data)
    assert response.status_code == 201
    doc = response.json()
    assert doc["status"] == "ready" and doc["file_type"] == "docx"
    text = owner.get(f"/api/clones/{clone_id}/documents/{doc['id']}").json()["extracted_text"]
    assert "I started with Python." in text and "Then Django and FastAPI." in text
    assert "Skill | FastAPI" in text  # tables are included


def test_scanned_pdf_fails_with_clear_message(owner: TestClient, clone_id: str):
    doc = upload(owner, clone_id, "scan.pdf", make_pdf(None)).json()
    assert doc["status"] == "failed"
    assert doc["error_message"] == "No extractable text found. OCR support is not enabled yet."


def test_corrupt_pdf_fails_without_crashing(owner: TestClient, clone_id: str):
    response = upload(owner, clone_id, "broken.pdf", b"%PDF-1.4\nthis is not really a pdf \x00\x01\x02")
    assert response.status_code == 201
    doc = response.json()
    assert doc["status"] == "failed"
    assert doc["error_message"] == "Could not read this PDF file. It may be damaged."
    assert "Traceback" not in response.text
    assert owner.get("/api/health").status_code == 200  # the app keeps running


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("virus.exe", b"MZ\x90\x00 executable"),
        ("script.sh", b"#!/bin/sh\nrm -rf /"),
        ("noextension", b"hello"),
        ("image.png", b"\x89PNG\r\n"),
        ("fake.pdf", b"MZ\x90\x00 an exe renamed to pdf"),  # extension alone isn't trusted
        ("fake.docx", b"just text, not a zip"),
        ("fake.txt", b"MZ\x90\x00\x00\x00 binary"),
        ("empty.txt", b""),
    ],
)
def test_unsupported_or_disguised_files_rejected(owner: TestClient, clone_id: str, storage_dir, name, data):
    response = upload(owner, clone_id, name, data)
    assert response.status_code == 400
    assert owner.get(f"/api/clones/{clone_id}/documents").json() == []
    assert stored_files(storage_dir) == []  # nothing left behind


def test_mismatched_client_mime_rejected(owner: TestClient, clone_id: str):
    response = upload(owner, clone_id, "cv.pdf", make_pdf("text"), "application/x-msdownload")
    assert response.status_code == 400


def test_oversized_file_rejected(owner: TestClient, clone_id: str, storage_dir, monkeypatch):
    monkeypatch.setattr(get_settings(), "max_upload_size_mb", 1)
    response = upload(owner, clone_id, "big.txt", b"a" * (1024 * 1024 + 200_000), "text/plain")
    assert response.status_code == 413
    assert response.json()["detail"] == "File is too large. Maximum allowed size is 1 MB."
    assert stored_files(storage_dir) == []


def test_file_just_over_limit_is_stopped_while_streaming(owner: TestClient, clone_id: str, storage_dir, monkeypatch):
    # Small enough to pass the Content-Length pre-check, still over the limit.
    monkeypatch.setattr(get_settings(), "max_upload_size_mb", 1)
    response = upload(owner, clone_id, "big.txt", b"a" * (1024 * 1024 + 10), "text/plain")
    assert response.status_code == 413
    assert stored_files(storage_dir) == []


def test_too_much_text_fails_cleanly(owner: TestClient, clone_id: str, monkeypatch):
    monkeypatch.setattr(get_settings(), "max_extracted_characters", 100)
    doc = upload(owner, clone_id, "long.txt", b"word " * 100, "text/plain").json()
    assert doc["status"] == "failed"
    assert "more than 100 characters" in doc["error_message"]


def test_path_traversal_filename_is_handled_safely(owner: TestClient, clone_id: str, storage_dir, db_session: Session):
    response = upload(owner, clone_id, "../../../etc/passwd.txt", b"harmless text content", "text/plain")
    assert response.status_code == 201
    doc = response.json()
    assert doc["original_filename"] == "passwd.txt"

    stored = db_session.get(Document, uuid.UUID(doc["id"]))
    files = stored_files(storage_dir)
    assert len(files) == 1
    assert files[0].name == stored.stored_filename != "passwd.txt"
    assert files[0].parent == storage_dir / "clones" / clone_id / "documents"

    windows = upload(owner, clone_id, "..\\..\\evil.txt", b"more harmless text", "text/plain").json()
    assert windows["original_filename"] == "evil.txt"


def test_unauthenticated_upload_rejected(make_client, clone_id: str):
    assert upload(make_client(), clone_id, "notes.txt", b"hello there", "text/plain").status_code == 401


# --- Manual knowledge -----------------------------------------------------------


def test_manual_knowledge_creation(owner: TestClient, clone_id: str):
    content = "I started learning Python in 2019 and later moved to FastAPI."
    response = owner.post(f"/api/clones/{clone_id}/knowledge", json={"title": "About My Career", "content": content})
    assert response.status_code == 201
    doc = response.json()
    assert (doc["source_type"], doc["status"], doc["character_count"]) == ("manual_text", "ready", len(content))
    assert doc["file_type"] is None and doc["original_filename"] is None
    detail = owner.get(f"/api/clones/{clone_id}/documents/{doc['id']}").json()
    assert detail["extracted_text"] == content


@pytest.mark.parametrize(
    "payload",
    [
        {"title": "Short", "content": "too short"},
        {"title": "", "content": "long enough content for knowledge"},
        {"title": "   ", "content": "long enough content for knowledge"},
        {"title": "Big", "content": "x" * 50_001},
        {"title": "Extra", "content": "long enough content for knowledge", "owner_id": str(uuid.uuid4())},
    ],
)
def test_manual_knowledge_validation(owner: TestClient, clone_id: str, payload: dict):
    assert owner.post(f"/api/clones/{clone_id}/knowledge", json=payload).status_code == 400


# --- List / get / delete / retry ------------------------------------------------


def test_list_documents_without_text(owner: TestClient, clone_id: str):
    upload(owner, clone_id, "a.txt", b"first document text", "text/plain")
    owner.post(f"/api/clones/{clone_id}/knowledge", json={"title": "Manual", "content": "manual knowledge text here"})
    listed = owner.get(f"/api/clones/{clone_id}/documents").json()
    assert [d["title"] for d in listed] == ["Manual", "a"]  # newest first
    assert all("extracted_text" not in d for d in listed)
    assert owner.get(f"/api/clones/{clone_id}").json()["knowledge_count"] == 2
    assert owner.get("/api/clones").json()[0]["knowledge_count"] == 2


def test_get_document(owner: TestClient, clone_id: str):
    doc = upload(owner, clone_id, "a.txt", b"<script>alert(1)</script> stays plain text", "text/plain").json()
    detail = owner.get(f"/api/clones/{clone_id}/documents/{doc['id']}")
    assert detail.status_code == 200
    assert detail.json()["extracted_text"] == "<script>alert(1)</script> stays plain text"


def test_get_document_of_another_clone_is_404(owner: TestClient, clone_id: str):
    second = owner.post("/api/clones", json={**CLONE, "name": "Second"}).json()["id"]
    doc = upload(owner, clone_id, "a.txt", b"belongs to the first clone", "text/plain").json()
    assert owner.get(f"/api/clones/{second}/documents/{doc['id']}").status_code == 404


def test_delete_document_removes_record_file_and_vectors(owner: TestClient, clone_id: str, storage_dir, qdrant):
    doc = upload(owner, clone_id, "a.txt", b"to be deleted soon", "text/plain").json()
    assert len(stored_files(storage_dir)) == 1
    assert owner.delete(f"/api/clones/{clone_id}/documents/{doc['id']}").status_code == 204
    assert owner.get(f"/api/clones/{clone_id}/documents").json() == []
    assert stored_files(storage_dir) == []
    assert qdrant.count_chunks(clone_id=uuid.UUID(clone_id), document_id=uuid.UUID(doc["id"])) == 0
    assert owner.delete(f"/api/clones/{clone_id}/documents/{doc['id']}").status_code == 404


def test_deleting_a_clone_removes_its_files(owner: TestClient, clone_id: str, storage_dir):
    upload(owner, clone_id, "a.txt", b"clone file content", "text/plain")
    assert owner.delete(f"/api/clones/{clone_id}").status_code == 204
    assert stored_files(storage_dir) == []


def test_retry_failed_document(owner: TestClient, clone_id: str, monkeypatch):
    from app.services import ingestion_service
    from app.services.document_extractor import ExtractionError

    real_extract = ingestion_service.extract_text

    def broken(*args, **kwargs):
        raise ExtractionError("Temporary failure")

    monkeypatch.setattr(ingestion_service, "extract_text", broken)
    doc = upload(owner, clone_id, "cv.pdf", make_pdf("Retry works"), "application/pdf").json()
    assert doc["status"] == "failed" and doc["error_message"] == "Temporary failure"

    monkeypatch.setattr(ingestion_service, "extract_text", real_extract)
    retried = owner.post(f"/api/clones/{clone_id}/documents/{doc['id']}/retry")
    assert retried.status_code == 200
    body = retried.json()
    assert body["id"] == doc["id"]  # same record, not a duplicate
    assert body["status"] == "ready" and body["error_message"] is None and body["character_count"] > 0
    assert len(owner.get(f"/api/clones/{clone_id}/documents").json()) == 1


def test_retry_only_for_failed_file_documents(owner: TestClient, clone_id: str):
    ready = upload(owner, clone_id, "a.txt", b"already ready text", "text/plain").json()
    manual = owner.post(f"/api/clones/{clone_id}/knowledge", json={"title": "M", "content": "manual knowledge content"}).json()
    assert owner.post(f"/api/clones/{clone_id}/documents/{ready['id']}/retry").status_code == 400
    assert owner.post(f"/api/clones/{clone_id}/documents/{manual['id']}/retry").status_code == 400


# --- Isolation between users ---------------------------------------------------


def test_knowledge_is_isolated_between_users(owner: TestClient, clone_id: str, other: TestClient):
    doc = upload(owner, clone_id, "cv.txt", b"User A private CV", "text/plain").json()
    base = f"/api/clones/{clone_id}/documents"

    # User B: every operation on User A's private clone looks like it doesn't exist.
    assert upload(other, clone_id, "x.txt", b"injected by user B", "text/plain").status_code == 404
    assert other.post(f"/api/clones/{clone_id}/knowledge", json={"title": "x", "content": "injected by user B!!"}).status_code == 404
    assert other.get(base).status_code == 404
    assert other.get(f"{base}/{doc['id']}").status_code == 404
    assert other.delete(f"{base}/{doc['id']}").status_code == 404
    assert other.post(f"{base}/{doc['id']}/retry").status_code == 404

    # Even with their own clone id in the URL, B can't reach A's document.
    b_clone = other.post("/api/clones", json={**CLONE, "name": "Sara AI"}).json()["id"]
    assert other.get(f"/api/clones/{b_clone}/documents/{doc['id']}").status_code == 404
    assert other.delete(f"/api/clones/{b_clone}/documents/{doc['id']}").status_code == 404

    # A public clone: B is recognised but still refused.
    owner.patch(f"/api/clones/{clone_id}", json={"visibility": "public"})
    assert other.get(f"{base}/{doc['id']}").status_code == 403
    assert other.delete(f"{base}/{doc['id']}").status_code == 403

    # User A's document is untouched.
    assert owner.get(f"{base}/{doc['id']}").json()["extracted_text"] == "User A private CV"
    assert len(owner.get(base).json()) == 1


# --- Text normalisation ----------------------------------------------------------


def test_normalize_text_keeps_paragraphs():
    raw = "Line  one\t\twith   spaces  \r\n\r\n\r\n\r\n  Paragraph two here \n"
    assert normalize_text(raw) == "Line one with spaces\n\nParagraph two here"
